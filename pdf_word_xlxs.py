import openpyxl
import streamlit as st
from PyPDF2 import PdfReader
from langchain.text_splitter import RecursiveCharacterTextSplitter
import os
from langchain_google_genai import GoogleGenerativeAIEmbeddings
import google.generativeai as genai
from langchain_community.vectorstores import FAISS
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain.chains.question_answering import load_qa_chain
from langchain.prompts import PromptTemplate
from dotenv import load_dotenv


load_dotenv()
genai.configure(api_key=os.getenv('GOOGLE_API_KEY'))


def get_pdf_text(path):
    text = ""
    for pdf in path:
        pr = PdfReader(pdf)
        for pg in pr.pages:
            text += pg.extract_text()
    return text


def get_text_chunks(text):
    ts = RecursiveCharacterTextSplitter(chunk_size=10000, chunk_overlap=1000)
    chunks = ts.split_text(text)
    return chunks


def get_vector_store(text_chunks):
    emb = GoogleGenerativeAIEmbeddings(model='models/embedding-001')
    vs = FAISS.from_texts(text_chunks, embedding=emb)
    #vs.save_local(f'C:/Users/ishwarya.thirumuruga/Desktop/Ishwarya/Pdf chat/uploaded')
    vs.save_local(f'vectors_storage')


def get_conversational_chain():
    prmt_temp = """
    Answer the question as detailed as possible from the provided context. 
    If the answer is not found in the context, just say "Unable to find the answer", don't provide a wrong answer.
    context:\n{context}\n
    Question:\n{question}\n
    Answer:
    """
    model = ChatGoogleGenerativeAI(model="gemini-pro", temperature=0.3)
    prompt = PromptTemplate(template=prmt_temp, input_variables=["context", "question"])
    chain = load_qa_chain(model, chain_type="stuff", prompt=prompt)
    return chain


def user_input(ques):
    emb = GoogleGenerativeAIEmbeddings(model='models/embedding-001')
    #new_db = FAISS.load_local(f'C:/Users/ishwarya.thirumuruga/Desktop/Ishwarya/Pdf chat/uploaded', emb, allow_dangerous_deserialization=True)
    new_db = FAISS.load_local(f'vectors_storage', emb, allow_dangerous_deserialization=True)
    dox = new_db.similarity_search(ques)
    chain = get_conversational_chain()
    res = chain.invoke(
        {"input_documents": dox, "question": ques},
    )
    return res["output_text"]


import tkinter as tk
from tkinter import filedialog
import PyPDF2
import docx
import docx2pdf
import builtins
from docx import Document
from openpyxl import workbook


def get_docx_text(path):
    text = ""
    doc = docx.Document(path)
    for para in doc.paragraphs:
        text += para.text
    return text


def get_excel_text(path):
    text = ""
    wb = openpyxl.load_workbook(path)
    for sheet in wb:
        for row in sheet.iter_rows(values_only=True):
            text += '\n'.join(map(str, row))
    return text




mas_read=""
fname=[]
def open_pdf():
    global mas_read 
    global fname
    file_paths = filedialog.askopenfilenames(title="Select a PDF", filetypes=[("PDF Files", "*.pdf"),("Word Documents","*.docx"),("Excel Files", "*.xlsx")])
    for file_path in file_paths:
        if file_path:
            if file_path.endswith('.pdf'):
                pdf_content=get_pdf_text([file_path])
                mas_read+=' '+pdf_content
                chunks = get_text_chunks(mas_read)
                get_vector_store(chunks)
                fname.append(file_path.split('/')[-1])
                #print(fname)
            elif file_path.endswith('.docx'):
                docx_content = get_docx_text(file_path)
                mas_read += ' '+docx_content
                chunks = get_text_chunks(mas_read)
                get_vector_store(chunks)
                fname.append(file_path.split('/')[-1])
                #print(fname)
            elif file_path.endswith('.xlsx'):
                excel_content = get_excel_text(file_path)
                mas_read += ' '+excel_content
                chunks = get_text_chunks(mas_read)
                get_vector_store(chunks)
                fname.append(file_path.split('/')[-1])
                print(fname)
        uploaded_label.config(text="\n".join(fname))
    print(mas_read)


def submit_question_answer():
    question = question_entry.get()
    ans=user_input(question)
    answer_text.insert(tk.INSERT,question+"?\n"+ans+"\n")
    # print("Question:", question)
    # print("Answer:", ans)
    question_entry.delete(0, tk.END)
    answer_text.see(tk.END)


root = tk.Tk()
root.title("PDF Viewer and Question/Answer")
root.geometry("800x600")


button = tk.Button(root, text="Open PDF", command=open_pdf)
button.pack(pady=10)


uploaded_label = tk.Label(root, text="")
uploaded_label.pack(pady=5)


question_label = tk.Label(root, text="Question:")
question_label.pack(pady=5)


question_entry = tk.Entry(root, width=50)
question_entry.pack(pady=5)


submit_button = tk.Button(root, text="Submit", command=submit_question_answer)
submit_button.pack(pady=10)


answer_label = tk.Label(root, text="Answer:")
answer_label.pack(pady=5)


answer_text = tk.Text(root, width=80, height=5)
answer_text.pack(pady=5)


root.mainloop()
